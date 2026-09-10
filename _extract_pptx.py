# -*- coding: utf-8 -*-
"""提取 PPTX 全部文本（含表格、备注）到文件，供分析考核需求"""
import io
from pptx import Presentation

SRC = r"c:\Users\罗国繁\Desktop\work-2\实习生AI Coding全栈开发实战考核.pptx"
OUT = r"c:\Users\罗国繁\Desktop\work-2\_pptx_text.txt"

prs = Presentation(SRC)
buf = io.StringIO()
for i, slide in enumerate(prs.slides, 1):
    buf.write(f"\n===== Slide {i} =====\n")
    # 普通形状中的文本
    for shape in slide.shapes:
        if shape.has_text_frame:
            txt = shape.text_frame.text.strip()
            if txt:
                buf.write(f"[shape:{shape.shape_type}] {txt}\n")
        if shape.has_table:
            buf.write("[table]\n")
            for row in shape.table.rows:
                cells = [c.text.strip() for c in row.cells]
                buf.write("  | " + " | ".join(cells) + "\n")
    # 备注
    if slide.has_notes_slide:
        notes = slide.notes_slide.notes_text_frame.text.strip()
        if notes:
            buf.write(f"[notes] {notes}\n")

with open(OUT, "w", encoding="utf-8") as f:
    f.write(buf.getvalue())
print("saved to", OUT)
